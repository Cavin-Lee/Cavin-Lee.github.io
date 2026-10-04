#!/usr/bin/env python3
"""Local-only editor for this static GitHub Pages repository."""

from __future__ import annotations

import copy
import hashlib
import html
import json
import mimetypes
import os
import re
import secrets
import shutil
import subprocess
import tempfile
import threading
import unicodedata
import urllib.parse
import webbrowser
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from gbt_citations import format_gbt


ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "data.js"
TOKEN = secrets.token_urlsafe(32)
HOST = "127.0.0.1"
PORT = int(os.environ.get("SITE_MANAGER_PORT", "8787"))
MAX_PDF_BYTES = 5_000_000
UPLOAD_KINDS = {"papers", "awards", "projects", "service"}
UPLOAD_CATEGORY_KIND = {"publications": "papers", "projects": "projects", "news": "awards", "personalAwards": "awards", "studentAwards": "awards", "service": "service"}
RECORD_CATEGORIES = {"education", "publications", "projects", "news", "personalAwards", "studentAwards", "service"}
LABELS = {"education": "教育经历", "publications": "论文", "projects": "科研项目", "news": "学术动态", "personalAwards": "个人荣誉", "studentAwards": "学生获奖", "service": "学术服务"}


class ManagerError(Exception):
    pass


def run_git(*args: str, check: bool = True) -> subprocess.CompletedProcess:
    result = subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True)
    if check and result.returncode:
        raise ManagerError((result.stderr or result.stdout).strip() or "Git 操作失败")
    return result


def atomic_bytes(path: Path, content: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=".site-manager-", dir=path.parent)
    try:
        with os.fdopen(fd, "wb") as stream:
            stream.write(content)
        os.replace(name, path)
    finally:
        if os.path.exists(name):
            os.unlink(name)


def data_bytes(data: dict) -> bytes:
    return ("window.SITE_DATA = " + json.dumps(data, ensure_ascii=False, indent=2) + ";\n").encode("utf-8")


def read_data() -> tuple[dict, str]:
    raw = DATA_PATH.read_bytes()
    text = raw.decode("utf-8")
    prefix = "window.SITE_DATA = "
    if not text.startswith(prefix) or not text.rstrip().endswith(";"):
        raise ManagerError("data.js 格式与管理工具不兼容")
    return json.loads(text[len(prefix):].rstrip().removesuffix(";")), hashlib.sha256(raw).hexdigest()


def bilingual(value, label: str, required: bool = True) -> dict:
    if not isinstance(value, dict):
        raise ManagerError(f"{label}需要填写中英文")
    result = {key: str(value.get(key, "")).strip() for key in ("zh", "en")}
    if required and not all(result.values()):
        raise ManagerError(f"{label}的中英文都不能留空")
    return result


def year(value, label="年份") -> int:
    try:
        number = int(value)
    except (ValueError, TypeError):
        raise ManagerError(f"{label}必须是年份") from None
    if not 1900 <= number <= 2100:
        raise ManagerError(f"{label}超出合理范围")
    return number


def pdf_path(value: str, optional: bool = True) -> str:
    value = str(value or "").strip()
    if not value and optional:
        return ""
    path = Path(value)
    if not value.startswith("assets/") or path.suffix.lower() != ".pdf" or ".." in path.parts or path.is_absolute():
        raise ManagerError(f"材料路径必须是 assets/ 下的 PDF：{value}")
    target = (ROOT / path).resolve()
    if not target.is_relative_to((ROOT / "assets").resolve()) or not target.is_file():
        raise ManagerError(f"找不到材料：{value}")
    validate_pdf(target)
    return value


def validate_pdf(path: Path) -> None:
    if path.stat().st_size >= MAX_PDF_BYTES:
        raise ManagerError(f"PDF 必须小于 5 MB：{path.name}")
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted or not reader.pages:
            raise ManagerError("PDF 不可加密且至少有一页")
        for index, page in enumerate(reader.pages, 1):
            width, height = float(page.mediabox.width), float(page.mediabox.height)
            short, long = sorted((width, height))
            if abs(short - 595.28) > 1.5 or abs(long - 841.89) > 1.5:
                raise ManagerError(f"第 {index} 页不是 A4：{path.name}")
    except ManagerError:
        raise
    except Exception as error:
        raise ManagerError(f"无法读取 PDF：{error}") from error


def clean_files(value, project=False, required=False) -> list[dict]:
    if not isinstance(value, list):
        raise ManagerError("材料列表格式不正确")
    result = []
    for file in value:
        if not isinstance(file, dict):
            raise ManagerError("材料条目格式不正确")
        item = {"path": pdf_path(file.get("path"), optional=False)}
        if project:
            item["label"] = bilingual(file.get("label"), "材料名称")
        result.append(item)
    if required and not result:
        raise ManagerError("获奖记录至少需要一份证明 PDF")
    return result


def validate_record(category: str, record: dict, data: dict) -> dict:
    if not isinstance(record, dict):
        raise ManagerError("条目格式不正确")
    if category == "education":
        return {"years": str(record.get("years", "")).strip(), "degree": bilingual(record.get("degree"), "学位"), "school": bilingual(record.get("school"), "学校")}
    if category == "publications":
        title = str(record.get("title", "")).strip()
        citation = str(record.get("citation", "")).strip()
        scholar = str(record.get("scholar", "")).strip()
        area = str(record.get("area", "")).strip()
        parsed = urllib.parse.urlsplit(scholar)
        params = urllib.parse.parse_qs(parsed.query)
        if not title or not citation or area not in data["areas"] or parsed.scheme != "https" or parsed.hostname != "scholar.google.com" or parsed.path != "/citations" or params.get("user") != ["XEfV8mkAAAAJ"] or not params.get("citation_for_view"):
            raise ManagerError("论文需填写标题、引文、研究方向和本人的 Google Scholar 记录链接")
        result = {"title": title, "year": year(record.get("year")), "citation": citation, "kind": "Google Scholar", "area": area, "scholar": scholar}
        result["gbtCitation"] = str(record.get("gbtCitation") or "").strip() or format_gbt(result)
        if record.get("file"):
            result["file"] = pdf_path(record["file"])
        return result
    if category == "projects":
        result = {key: bilingual(record.get(key), key) for key in ("title", "funder", "role")}
        if record.get("note") and any(str(x).strip() for x in record["note"].values()):
            result["note"] = bilingual(record["note"], "项目说明")
        files = clean_files(record.get("files", []), project=True)
        if files:
            result["files"] = files
        return result
    if category in ("personalAwards", "studentAwards"):
        result = {"year": year(record.get("year")), "title": bilingual(record.get("title"), "奖项名称"), "files": clean_files(record.get("files", []), required=True)}
        if record.get("note") and any(str(x).strip() for x in record["note"].values()):
            result["note"] = bilingual(record["note"], "获奖说明")
        return result
    if category == "news":
        result = {"year": year(record.get("year")), "title": bilingual(record.get("title"), "新闻标题"), "detail": bilingual(record.get("detail"), "新闻正文")}
        if record.get("file"):
            result["file"] = pdf_path(record["file"])
        return result
    if category == "service":
        result = {"name": bilingual(record.get("name"), "任职名称")}
        if record.get("file"):
            result["file"] = pdf_path(record["file"])
        return result
    raise ManagerError("不支持的内容类别")


def record_label(category: str, record: dict) -> str:
    value = record.get("title") or record.get("name") or record.get("degree") or record.get("years") or "条目"
    if isinstance(value, dict):
        value = value.get("zh") or value.get("en") or "条目"
    return str(value).replace("\n", " ")[:100]


def md_text(value: str) -> str:
    return str(value).replace("\n", " ").replace("|", "\\|").strip()


def content_index(data: dict) -> str:
    lines = ["# 网站内容索引", "", "此文件由本地管理工具根据 `data.js` 自动生成，供后续智能体检索。请在管理页面修改内容，不要直接改此文件。", "", f"生成时间：{datetime.now().astimezone().isoformat(timespec='seconds')}", ""]
    profile = data.get("profile") or profile_from_pages()
    lines += ["## 主页文案", "", f"- 职称：{md_text(profile['role']['zh'])} / {md_text(profile['role']['en'])}", "- 学术经历："]
    lines += [f"  - {md_text(paragraph)}" for paragraph in profile["about"]["zh"]]
    lines += ["- 研究方向："] + [f"  - [{md_text(topic['title']['zh'])}](cn/#papers-{key}) / [{md_text(topic['title']['en'])}](index.html#papers-{key})" for key, topic in zip(data["areas"], profile["topics"])] + [""]
    for category in ("news", "education", "projects", "service", "personalAwards", "studentAwards", "publications"):
        entries = [(g, x) for g, group in enumerate(data["service"]) for x in group["items"]] if category == "service" else [(None, x) for x in data[category]]
        lines += [f"## {LABELS[category]}（{len(entries)}）", ""]
        for number, (group, item) in enumerate(entries, 1):
            name = item.get("title") or item.get("name") or item.get("degree") or item.get("years") or "条目"
            label = name.get("zh", "") if isinstance(name, dict) else str(name)
            prefix = f"{item['year']} · " if item.get("year") else ""
            group_label = f"[{data['service'][group]['title']['zh']}] " if group is not None else ""
            sources = []
            if item.get("scholar"):
                sources.append(f"[Scholar]({item['scholar']})")
            for file in item.get("files", []):
                sources.append(f"[PDF]({file['path']})")
            if item.get("file"):
                sources.append(f"[PDF]({item['file']})")
            if category == "education":
                label = f"{item['years']} · {item['degree']['zh']} · {item['school']['zh']}"
            lines.append(f"{number}. {prefix}{group_label}{md_text(label)}" + (" · " + " · ".join(sources) if sources else ""))
            if isinstance(name, dict) and name.get("en"):
                lines.append(f"   - English: {md_text(name['en'])}")
            if category == "publications":
                lines.append(f"   - {md_text(item['citation'])} · 方向：{md_text(data['areas'][item['area']]['zh'])}")
                lines.append(f"   - 简历引文（GB/T 7714—2025）：{md_text(item.get('gbtCitation') or format_gbt(item))}")
            elif category == "projects":
                lines.append(f"   - {md_text(item['funder']['zh'])} · {md_text(item['role']['zh'])}")
            elif category == "news":
                lines.append(f"   - {md_text(item['detail']['zh'])}")
            elif category == "education":
                lines.append(f"   - {md_text(item['degree']['en'])} · {md_text(item['school']['en'])}")
            if item.get("note"):
                lines.append(f"   - 说明：{md_text(item['note']['zh'])}")
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def log_entry(action: str, category: str, label: str, note: str = "") -> str:
    stamp = datetime.now().astimezone().isoformat(timespec="seconds")
    safe_note = md_text(note)[:300]
    return f"\n## {stamp}\n\n- 操作：{action}\n- 类别：{category}\n- 条目：{md_text(label)}\n" + (f"- 备注：{safe_note}\n" if safe_note else "")


def missing_publications(data: dict) -> list[dict]:
    return sorted((x for x in data["publications"] if not x.get("file")), key=lambda x: (-(x.get("year") or 0), x["title"]))


def write_missing_reports(data: dict, temp: Path) -> dict[str, Path]:
    missing = missing_publications(data)
    md = ["# 缺少本地全文的论文", "", "以 Google Scholar 个人主页记录为准。本清单仅表示当前仓库未收录可确认匹配的全文。", "", f"共 {len(missing)} 篇待补。", ""]
    md += [f"- {p.get('year') or '未标年份'} · {p['title']} · {p['scholar']}" for p in missing]
    markdown = temp / "missing-papers.md"
    markdown.write_text("\n".join(md) + "\n", encoding="utf-8")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "待补论文"
    sheet.append(["序号", "研究方向", "年份", "论文题目", "发表信息", "Google Scholar 记录", "状态"])
    for number, item in enumerate(missing, 1):
        sheet.append([number, data["areas"][item["area"]]["zh"], item.get("year") or "", item["title"], item["citation"], item["scholar"], "待补全文"])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 8, "B": 24, "C": 10, "D": 65, "E": 80, "F": 55, "G": 14}.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17323A")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.row_dimensions[1].height = 25
    xlsx = temp / "missing-papers.xlsx"
    workbook.save(xlsx)
    pdf = temp / "missing-papers.pdf"
    write_pdf(pdf, "zh", "缺少本地全文的论文清单", [("待补论文", [f"{p.get('year') or '未标年份'} · {p['title']}" for p in missing])])
    return {"missing-papers.md": markdown, "assets/missing-papers.xlsx": xlsx, "assets/missing-papers.pdf": pdf}


pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))


def write_pdf(path: Path, lang: str, title: str, sections: list[tuple[str, list[str]]], numbered: bool = False, intro: list[str] | None = None, portrait: Path | None = None) -> None:
    font = "STSong-Light" if lang == "zh" else "Helvetica"
    title_style = ParagraphStyle("title", fontName=font, fontSize=18, leading=25, textColor=colors.HexColor("#202329"), spaceAfter=12, wordWrap="CJK")
    section_style = ParagraphStyle("section", fontName=font, fontSize=11.5, leading=17, textColor=colors.HexColor("#163a78"), spaceBefore=16, spaceAfter=7, wordWrap="CJK")
    body_style = ParagraphStyle("body", fontName=font, fontSize=8.7, leading=13.3, textColor=colors.HexColor("#202329"), leftIndent=12, firstLineIndent=-12, spaceAfter=5, wordWrap="CJK")
    if portrait:
        intro_style = ParagraphStyle("intro", fontName=font, fontSize=8.7, leading=13.3, textColor=colors.HexColor("#202329"), spaceAfter=5, wordWrap="CJK")
        header_text = [Paragraph(escape(title), title_style)] + [Paragraph(escape(line.replace("·", " / ")), intro_style) for line in intro or [] if line]
        ratio = 1929 / 1279
        photo = Image(str(portrait), width=72, height=72 * ratio)
        photo.hAlign = "RIGHT"
        header = Table([[header_text, photo]], colWidths=[419, 92])
        header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0), ("LINEBELOW", (0, 0), (-1, -1), 1, colors.HexColor("#a12f35"))]))
        story = [header, Spacer(1, 10)]
    else:
        story = [Paragraph(escape(title), title_style)]
        for line in intro or []:
            story.append(Paragraph(escape(line), body_style))
    for heading, entries in sections:
        story.append(Paragraph(escape(heading), section_style))
        for number, entry in enumerate(entries, 1):
            label = f"[{number}]" if heading in ("论文成果", "Publications") else f"{number}."
            marker = f'<font color="#a12f35">{label}</font> ' if numbered else "- "
            story.append(Paragraph(marker + escape(str(entry).replace("·", " / ")), body_style))
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=55, title=title, author="Wei-Kai Li")
    doc.build(story)
    validate_pdf(path)


def write_cv_reports(data: dict, temp: Path) -> dict[str, Path]:
    output = {}
    profile = data.get("profile") or profile_from_pages()
    for lang in ("zh", "en"):
        zh = lang == "zh"
        title = "李伟凯 学术简历" if zh else "Wei-Kai Li Academic CV"
        sections = [
            (("教育经历" if zh else "Education"), [f"{x['years']} · {x['degree'][lang]} · {x['school'][lang]}" for x in data["education"]]),
            (("科研项目" if zh else "Research Projects"), [f"{x['title'][lang]} · {x['funder'][lang]} · {x['role'][lang]}" for x in data["projects"]]),
            (("论文成果" if zh else "Publications"), [x.get("gbtCitation") or format_gbt(x) for x in sorted(data["publications"], key=lambda p: -(p.get("year") or 0))]),
            (("学术服务" if zh else "Academic Service"), [item["name"][lang] for group in data["service"] for item in group["items"]]),
            (("个人荣誉" if zh else "Personal Recognition"), [f"{x['year']} · {x['title'][lang]}" for x in data["personalAwards"]]),
            (("指导学生获奖" if zh else "Student Awards"), [f"{x['year']} · {x['title'][lang]}" for x in data["studentAwards"]]),
        ]
        intro = [profile.get("role", {}).get(lang, ""), "leeweikai@sdjzu.edu.cn · Google Scholar: scholar.google.com/citations?user=XEfV8mkAAAAJ", " / ".join(profile.get("affiliations", {}).get(lang, []))]
        path = temp / f"cv-{lang}.pdf"
        write_pdf(path, lang, title, sections, numbered=True, intro=intro, portrait=ROOT / "assets/profile.jpeg")
        output[f"assets/cv-{lang}.pdf"] = path
    return output


def write_project_report(data: dict, temp: Path) -> dict[str, Path]:
    path = temp / "projects.pdf"
    sections = [("主持项目", [f"{x['title']['zh']} · {x['funder']['zh']}" for x in data["projects"] if x["role"]["zh"] == "主持"]), ("参与项目", [f"{x['title']['zh']} · {x['funder']['zh']}" for x in data["projects"] if x["role"]["zh"] != "主持"])]
    write_pdf(path, "zh", "李伟凯 科研项目清单", sections)
    return {"assets/projects.pdf": path}


def profile_from_pages() -> dict:
    parts = {}
    for lang, page in (("en", ROOT / "index.html"), ("zh", ROOT / "cn/index.html")):
        source = page.read_text(encoding="utf-8")
        def one(pattern):
            match = re.search(pattern, source, re.S)
            if not match:
                raise ManagerError(f"无法识别 {page.name} 的主页文案")
            return html.unescape(re.sub(r"<[^>]+>", "", match.group(1))).strip()
        role = one(r'<p class="hero-role">(.*?)</p>')
        aff = re.search(r'<p class="hero-affiliation">(.*?)</p>', source, re.S).group(1)
        affiliations = [html.unescape(x.strip()) for x in aff.split("<br>")]
        summary = one(r'<p class="hero-summary">(.*?)</p>')
        body = re.search(r'<div class="body-copy">(.*?)</div>', source, re.S).group(1)
        about = [html.unescape(re.sub(r"<[^>]+>", "", x)).strip() for x in re.findall(r"<p>(.*?)</p>", body, re.S)]
        grid = re.search(r'<div class="topic-grid">(.*?)</div>', source, re.S).group(1)
        topics = [(html.unescape(a), html.unescape(b)) for a, b in re.findall(r"<h3>(.*?)</h3><p>(.*?)</p>", grid, re.S)]
        parts[lang] = {"role": role, "affiliations": affiliations, "summary": summary, "about": about, "topics": topics}
    if len(parts["zh"]["topics"]) != len(parts["en"]["topics"]):
        raise ManagerError("中英文研究方向数量不一致")
    return {"role": {lang: parts[lang]["role"] for lang in parts}, "affiliations": {lang: parts[lang]["affiliations"] for lang in parts}, "summary": {lang: parts[lang]["summary"] for lang in parts}, "about": {lang: parts[lang]["about"] for lang in parts}, "topics": [{"title": {lang: parts[lang]["topics"][i][0] for lang in parts}, "description": {lang: parts[lang]["topics"][i][1] for lang in parts}} for i in range(len(parts["zh"]["topics"]))]}


def validate_profile(value: dict, data: dict) -> dict:
    if not isinstance(value, dict):
        raise ManagerError("主页文案格式不正确")
    result = {"role": bilingual(value.get("role"), "职称"), "summary": bilingual(value.get("summary"), "简介")}
    for field in ("affiliations", "about"):
        obj = value.get(field)
        if not isinstance(obj, dict):
            raise ManagerError(f"{field}缺少中英文")
        result[field] = {}
        for lang in ("zh", "en"):
            lines = obj.get(lang)
            if not isinstance(lines, list) or not lines or not all(str(x).strip() for x in lines):
                raise ManagerError(f"{field} 的 {lang} 内容不能为空")
            result[field][lang] = [str(x).strip() for x in lines]
    topics = value.get("topics")
    if not isinstance(topics, list) or not 1 <= len(topics) <= 12:
        raise ManagerError("研究方向需保留 1 至 12 项")
    result["topics"] = [{"title": bilingual(x.get("title"), "研究方向名称"), "description": bilingual(x.get("description"), "研究方向说明")} for x in topics]
    expected = list(data["areas"].values())
    if len(result["topics"]) != len(expected) or any(topic["title"] != area for topic, area in zip(result["topics"], expected)):
        raise ManagerError("研究方向名称及顺序须与论文分类一致；可修改每项说明")
    return result


def render_profile_html(source: str, profile: dict, lang: str, areas: list[str]) -> str:
    def substitute(pattern: str, replacement: str, text: str) -> str:
        result, count = re.subn(pattern, lambda match: match.group(1) + replacement + match.group(3), text, count=1, flags=re.S)
        if count != 1:
            raise ManagerError("主页结构已变化，无法安全保存文案")
        return result
    esc = lambda value: html.escape(value, quote=False)
    source = substitute(r'(<p class="hero-role">)(.*?)(</p>)', esc(profile["role"][lang]), source)
    source = substitute(r'(<p class="hero-affiliation">)(.*?)(</p>)', "<br>".join(esc(x) for x in profile["affiliations"][lang]), source)
    source = substitute(r'(<p class="hero-summary">)(.*?)(</p>)', esc(profile["summary"][lang]), source)
    source = substitute(r'(<div class="body-copy">)(.*?)(</div>)', "".join(f"<p>{esc(x)}</p>" for x in profile["about"][lang]), source)
    link_label = "查看论文 ↗" if lang == "zh" else "View papers ↗"
    cards = "".join(f'<a class="topic-card" href="#papers-{key}" data-area="{key}"><span>{i:02d}</span><h3>{esc(x["title"][lang])}</h3><p>{esc(x["description"][lang])}</p><span class="topic-link">{link_label}</span></a>' for i, (key, x) in enumerate(zip(areas, profile["topics"]), 1))
    source = substitute(r'(<div class="topic-grid">)(.*?)(</div>)', cards, source)
    return source


def bump_data_version(source: str) -> str:
    return re.sub(r'(data\.js\?v=)[A-Za-z0-9]+', lambda match: match.group(1) + str(int(datetime.now().timestamp() * 1000)), source)


def save_changes(data: dict, action: str, category: str, label: str, note: str = "", profile_changed: bool = False) -> None:
    with tempfile.TemporaryDirectory(prefix="site-manager-") as directory:
        temp = Path(directory)
        staged: dict[str, bytes] = {"data.js": data_bytes(data)}
        staged["CONTENT_INDEX.md"] = content_index(data).encode("utf-8")
        old_log = (ROOT / "CONTENT_CHANGELOG.md").read_text(encoding="utf-8") if (ROOT / "CONTENT_CHANGELOG.md").exists() else "# 网站内容修改记录\n\n按本地保存时间倒序检索 Git 历史和下列记录。\n"
        staged["CONTENT_CHANGELOG.md"] = (old_log + log_entry(action, LABELS.get(category, category), label, note)).encode("utf-8")
        reports: dict[str, Path] = {}
        if category == "publications":
            reports.update(write_missing_reports(data, temp))
        if category == "projects":
            reports.update(write_project_report(data, temp))
        if category not in ("news", "upload"):
            reports.update(write_cv_reports(data, temp))
        for relative, path in reports.items():
            staged[relative] = path.read_bytes()
        for relative in ("index.html", "cn/index.html", "cv-builder/index.html"):
            source = (ROOT / relative).read_text(encoding="utf-8")
            if profile_changed and relative in ("index.html", "cn/index.html"):
                source = render_profile_html(source, data["profile"], "zh" if relative.startswith("cn/") else "en", list(data["areas"]))
            staged[relative] = bump_data_version(source).encode("utf-8")
        for relative, content in staged.items():
            atomic_bytes(ROOT / relative, content)


def mutate_record(payload: dict) -> dict:
    data, revision = read_data()
    if payload.get("revision") != revision:
        raise ManagerError("内容已被其他程序修改。请刷新管理页面后重试。")
    category = payload.get("category")
    action = payload.get("action")
    if category not in RECORD_CATEGORIES or action not in {"create", "update", "delete", "up", "down"}:
        raise ManagerError("不支持的编辑操作")
    group = payload.get("groupIndex", 0)
    if category == "service":
        if not isinstance(group, int) or not 0 <= group < len(data["service"]):
            raise ManagerError("学术服务分组无效")
        items = data["service"][group]["items"]
    else:
        items = data[category]
    index = payload.get("index", -1)
    if action != "create" and (not isinstance(index, int) or not 0 <= index < len(items)):
        raise ManagerError("所选条目已变化，请刷新页面")
    before = items[index] if action != "create" else None
    if action == "create":
        record = validate_record(category, payload.get("record"), data)
        if category == "publications" and any(x["scholar"] == record["scholar"] for x in items):
            raise ManagerError("这条 Google Scholar 论文记录已存在")
        items.insert(0, record)
    elif action == "update":
        submitted = payload.get("record")
        if category == "publications" and isinstance(submitted, dict):
            source_changed = any(str(submitted.get(key) or "") != str(before.get(key) or "") for key in ("title", "year", "citation"))
            if source_changed and submitted.get("gbtCitation") == before.get("gbtCitation"):
                submitted = dict(submitted, gbtCitation="")
        record = validate_record(category, submitted, data)
        if category == "publications" and any(i != index and x["scholar"] == record["scholar"] for i, x in enumerate(items)):
            raise ManagerError("这条 Google Scholar 论文记录已存在")
        items[index] = record
    elif action == "delete":
        items.pop(index)
        record = before
    else:
        destination = index + (-1 if action == "up" else 1)
        if not 0 <= destination < len(items):
            raise ManagerError("该条目已经位于边界")
        items[index], items[destination] = items[destination], items[index]
        record = before
    label = record_label(category, record)
    save_changes(data, action, category, label, str(payload.get("note", "")))
    return {"data": data, "revision": hashlib.sha256(DATA_PATH.read_bytes()).hexdigest(), "message": f"已在本地{ {'create':'新增','update':'保存','delete':'删除','up':'上移','down':'下移'}[action]}：{label}"}


def save_profile(payload: dict) -> dict:
    data, revision = read_data()
    if payload.get("revision") != revision:
        raise ManagerError("内容已变化。请刷新管理页面后重试。")
    data["profile"] = validate_profile(payload.get("profile"), data)
    save_changes(data, "更新", "主页文案", "中英文职称、单位、简介与研究方向", str(payload.get("note", "")), profile_changed=True)
    return {"data": data, "revision": hashlib.sha256(DATA_PATH.read_bytes()).hexdigest(), "message": "主页文案已保存在本地"}


def status() -> dict:
    output = run_git("status", "--porcelain=v1", "-z", "--untracked-files=all").stdout
    parts = output.split("\0")
    changes = []
    for part in parts:
        if not part:
            continue
        code, name = part[:2], part[3:]
        if code[0] in "RC" or code[1] in "RC":
            continue
        recommended = name in {"data.js", "app.js", "styles.css", "CONTENT_INDEX.md", "CONTENT_CHANGELOG.md", "index.html", "cn/index.html", "cv-builder/index.html", "cv-builder/builder.js", "cv-builder/builder.css", "missing-papers.md"} or name.startswith("assets/")
        changes.append({"path": name, "status": code.strip(), "recommended": recommended})
    ahead = int(run_git("rev-list", "--count", "origin/main..HEAD").stdout.strip() or "0")
    return {"changes": changes, "ahead": ahead, "branch": run_git("branch", "--show-current").stdout.strip()}


def referenced_pdfs(data: dict) -> set[str]:
    paths = set()
    for category in ("publications", "projects", "news", "personalAwards", "studentAwards"):
        for item in data[category]:
            if item.get("file"):
                paths.add(item["file"])
            paths.update(file["path"] for file in item.get("files", []))
    for group in data["service"]:
        paths.update(item["file"] for item in group["items"] if item.get("file"))
    return paths


def publish(payload: dict) -> dict:
    current = status()
    if current["branch"] != "main":
        raise ManagerError("请先切回 main 分支")
    if run_git("diff", "--cached", "--name-only").stdout.strip():
        raise ManagerError("Git 暂存区已有其他改动，请先处理后再用管理页面发布")
    paths = payload.get("paths", [])
    known = {entry["path"] for entry in current["changes"]}
    if not isinstance(paths, list) or any(path not in known for path in paths):
        raise ManagerError("待提交文件已变化，请刷新状态")
    references = referenced_pdfs(read_data()[0])
    generated = {"assets/cv-zh.pdf", "assets/cv-en.pdf", "assets/projects.pdf", "assets/missing-papers.pdf"}
    unattached = [path for path in paths if path.startswith("assets/") and path.endswith(".pdf") and path not in references and path not in generated]
    if unattached:
        raise ManagerError("这些 PDF 尚未关联到条目，请先保存对应条目：" + "、".join(unattached[:3]))
    run_git("fetch", "origin", "main")
    if run_git("merge-base", "--is-ancestor", "origin/main", "HEAD", check=False).returncode:
        raise ManagerError("GitHub 上已有新修改。请先同步仓库再发布。")
    if paths:
        try:
            run_git("add", "--", *paths)
            run_git("diff", "--cached", "--check")
            message = str(payload.get("message", "")).strip()[:120] or "Update academic website content"
            run_git("commit", "-m", message)
        except ManagerError:
            run_git("reset", "--", *paths, check=False)
            raise
    elif not current["ahead"]:
        raise ManagerError("目前没有需要发布的改动")
    run_git("push", "origin", "main")
    commit = run_git("rev-parse", "HEAD").stdout.strip()
    return {"message": "已提交并推送到 GitHub。GitHub Pages 通常稍后完成更新。", "commit": commit, "url": f"https://github.com/Cavin-Lee/Cavin-Lee.github.io/commit/{commit}"}


class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt: str, *args) -> None:
        print(f"[{self.log_date_time_string()}] {fmt % args}")

    def response(self, payload: dict, code=200) -> None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def auth(self) -> bool:
        origin = self.headers.get("Origin")
        expected = f"http://{HOST}:{PORT}"
        return self.headers.get("X-Manager-Token") == TOKEN and origin in (None, expected)

    def read_json(self) -> dict:
        length = int(self.headers.get("Content-Length", "0"))
        if not 0 < length < 2_000_000:
            raise ManagerError("请求内容过大或为空")
        value = json.loads(self.rfile.read(length))
        if not isinstance(value, dict):
            raise ManagerError("请求格式不正确")
        return value

    def serve_file(self, path: Path, content_type: str | None = None, substitute_token=False) -> None:
        if not path.is_file():
            self.send_error(404)
            return
        body = path.read_bytes()
        if substitute_token:
            body = body.replace(b"__LOCAL_MANAGER_TOKEN__", TOKEN.encode("ascii"))
        self.send_response(200)
        self.send_header("Content-Type", content_type or mimetypes.guess_type(path.name)[0] or "application/octet-stream")
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self) -> None:
        route = urllib.parse.urlsplit(self.path).path
        if route in ("/", "/manage", "/manage/"):
            self.serve_file(ROOT / "tools/manage.html", "text/html; charset=utf-8", substitute_token=True)
            return
        if route in ("/manage/manage.js", "/manage/manage.css"):
            self.serve_file(ROOT / "tools" / Path(route).name)
            return
        if route.startswith("/api/"):
            if not self.auth():
                self.response({"error": "仅允许本地管理页面访问"}, 403)
                return
            try:
                if route == "/api/state":
                    data, revision = read_data()
                    self.response({"data": data, "revision": revision, "profile": data.get("profile") or profile_from_pages(), "status": status()})
                elif route == "/api/status":
                    self.response(status())
                elif route == "/api/files":
                    files = [str(p.relative_to(ROOT)) for kind in sorted(UPLOAD_KINDS) for p in sorted((ROOT / "assets" / kind).glob("*.pdf"))]
                    self.response({"files": files})
                elif route == "/api/log":
                    self.response({"text": (ROOT / "CONTENT_CHANGELOG.md").read_text(encoding="utf-8") if (ROOT / "CONTENT_CHANGELOG.md").exists() else "尚无管理页面修改记录。"})
                else:
                    self.response({"error": "未知接口"}, 404)
            except (ManagerError, ValueError) as error:
                self.response({"error": str(error)}, 400)
            return
        if route.startswith("/site/"):
            relative = urllib.parse.unquote(route[len("/site/"):])
            candidate = (ROOT / relative).resolve()
            if not candidate.is_relative_to(ROOT) or ".git" in candidate.parts or ".venv" in candidate.parts or "tools" in candidate.parts:
                self.send_error(403)
                return
            if candidate.is_dir():
                candidate /= "index.html"
            self.serve_file(candidate)
            return
        self.send_error(404)

    def do_POST(self) -> None:
        route = urllib.parse.urlsplit(self.path).path
        if not self.auth():
            self.response({"error": "仅允许本地管理页面访问"}, 403)
            return
        try:
            if route == "/api/records":
                result = mutate_record(self.read_json())
            elif route == "/api/profile":
                result = save_profile(self.read_json())
            elif route == "/api/publish":
                result = publish(self.read_json())
            elif route == "/api/upload":
                category = self.headers.get("X-Record-Category", "")
                kind = UPLOAD_CATEGORY_KIND.get(category)
                if not kind:
                    raise ManagerError("请从对应的论文、项目、获奖或学术服务条目中上传 PDF")
                length = int(self.headers.get("Content-Length", "0"))
                if not 0 < length < MAX_PDF_BYTES:
                    raise ManagerError("只能上传小于 5 MB 的 PDF")
                name = urllib.parse.unquote(self.headers.get("X-File-Name", "document.pdf"))
                if not name.lower().endswith(".pdf"):
                    raise ManagerError("只能上传 PDF 文件")
                content = self.rfile.read(length)
                with tempfile.NamedTemporaryFile(suffix=".pdf") as temporary:
                    temporary.write(content)
                    temporary.flush()
                    validate_pdf(Path(temporary.name))
                stem = unicodedata.normalize("NFKD", Path(name).stem).encode("ascii", "ignore").decode("ascii")
                stem = re.sub(r"[^a-zA-Z0-9_-]+", "-", stem).strip("-").lower()[:55] or "document"
                target = ROOT / "assets" / kind / f"{stem}-{hashlib.sha256(content).hexdigest()[:8]}.pdf"
                if not target.exists():
                    atomic_bytes(target, content)
                    previous = (ROOT / "CONTENT_CHANGELOG.md").read_text(encoding="utf-8") if (ROOT / "CONTENT_CHANGELOG.md").exists() else "# 网站内容修改记录\n"
                    related = urllib.parse.unquote(self.headers.get("X-Record-Title", ""))[:100]
                    atomic_bytes(ROOT / "CONTENT_CHANGELOG.md", (previous + log_entry("上传 PDF", LABELS[category], related or str(target.relative_to(ROOT)), str(target.relative_to(ROOT)))).encode("utf-8"))
                result = {"path": str(target.relative_to(ROOT)), "message": "PDF 已加入当前条目。请保存条目后再发布。"}
            else:
                self.response({"error": "未知接口"}, 404)
                return
            self.response(result)
        except (ManagerError, ValueError, json.JSONDecodeError) as error:
            self.response({"error": str(error)}, 400)
        except Exception as error:
            self.response({"error": f"操作失败：{error}"}, 500)


def main() -> None:
    if not (ROOT / ".git").exists():
        raise SystemExit("请在 GitHub 仓库中运行此工具")
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    url = f"http://{HOST}:{PORT}/manage/"
    print(f"本地网站管理页面：{url}\n仅此电脑可访问；按 Ctrl+C 关闭。", flush=True)
    if not os.environ.get("SITE_MANAGER_NO_BROWSER"):
        threading.Timer(0.7, lambda: webbrowser.open(url)).start()
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\n管理页面已关闭。")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
